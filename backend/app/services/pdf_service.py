import uuid
from pathlib import Path
from typing import List, Dict, Any, Tuple
import pymupdf  # High-performance PyMuPDF
from pypdf import PdfReader
from langchain_text_splitters import RecursiveCharacterTextSplitter

class PDFService:
    def __init__(self, chunk_size: int = 800, chunk_overlap: int = 150):
        self.splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
            separators=["\n\n", "\n", " ", ""]
        )

    def process_pdf(self, file_path: Path, filename: str) -> Tuple[List[Dict[str, Any]], int, str]:
        """
        Reads a PDF file page by page using PyMuPDF (with pypdf fallback),
        extracts text, and chunks it while preserving page numbers.
        Returns (list_of_chunks_with_metadata, total_pages, doc_id).
        """
        chunks: List[Dict[str, Any]] = []
        doc_id = str(uuid.uuid4())
        total_pages = 0

        # 1. Primary extractor: PyMuPDF (fast & robust for complex fonts/layouts)
        try:
            doc = pymupdf.open(str(file_path))
            total_pages = len(doc)

            for page_idx in range(total_pages):
                page = doc[page_idx]
                text = page.get_text("text") or ""
                
                # If standard text mode was empty, try extracting text blocks
                if not text.strip():
                    blocks = page.get_text("blocks")
                    if blocks:
                        text = "\n".join([b[4] for b in blocks if len(b) > 4 and isinstance(b[4], str)])

                if not text.strip():
                    continue

                page_number = page_idx + 1
                page_chunks = self.splitter.split_text(text)

                for chunk_idx, chunk_text in enumerate(page_chunks):
                    chunks.append({
                        "chunk_id": f"{doc_id}_p{page_number}_c{chunk_idx}",
                        "doc_id": doc_id,
                        "filename": filename,
                        "page_number": page_number,
                        "text": chunk_text
                    })
            doc.close()
        except Exception as e:
            print(f"PyMuPDF error: {e}, falling back to pypdf...")

        # 2. Fallback extractor: pypdf (if PyMuPDF found 0 chunks)
        if not chunks:
            try:
                reader = PdfReader(str(file_path))
                total_pages = len(reader.pages)
                for page_idx, page in enumerate(reader.pages):
                    text = page.extract_text() or ""
                    if not text.strip():
                        continue

                    page_number = page_idx + 1
                    page_chunks = self.splitter.split_text(text)

                    for chunk_idx, chunk_text in enumerate(page_chunks):
                        chunks.append({
                            "chunk_id": f"{doc_id}_p{page_number}_c{chunk_idx}",
                            "doc_id": doc_id,
                            "filename": filename,
                            "page_number": page_number,
                            "text": chunk_text
                        })
            except Exception as e:
                print(f"pypdf fallback error: {e}")

        # 3. Intelligent Fallback: Gemini Multimodal AI OCR for Scanned / Image PDFs
        if not chunks:
            try:
                print(f"[PDFService] No digital text in '{filename}'. Activating Gemini Multimodal OCR fallback...")
                from app.services.ai_service import ai_service
                transcribed_text = ai_service.ocr_scanned_pdf(file_path)
                if transcribed_text and transcribed_text.strip():
                    if total_pages == 0:
                        total_pages = max(1, (len(transcribed_text) // 2000) + 1)

                    ocr_chunks = self.splitter.split_text(transcribed_text)
                    total_chunks = max(1, len(ocr_chunks))
                    for chunk_idx, chunk_text in enumerate(ocr_chunks):
                        page_num = min(total_pages, (chunk_idx * total_pages // total_chunks) + 1)
                        chunks.append({
                            "chunk_id": f"{doc_id}_p{page_num}_c{chunk_idx}",
                            "doc_id": doc_id,
                            "filename": filename,
                            "page_number": page_num,
                            "text": chunk_text
                        })
                    print(f"[PDFService] AI OCR recovered {len(chunks)} clauses across ~{total_pages} page(s).")
            except Exception as ocr_err:
                print(f"[PDFService] AI OCR fallback failed for {filename}: {ocr_err}")

        return chunks, total_pages, doc_id

    def process_docx(self, file_path: Path, filename: str) -> Tuple[List[Dict[str, Any]], int, str]:
        """
        Reads a Microsoft Word (.docx) document, extracts paragraphs and table text,
        and chunks content while approximating page distribution.
        Returns (list_of_chunks_with_metadata, estimated_pages, doc_id).
        """
        import docx
        chunks: List[Dict[str, Any]] = []
        doc_id = str(uuid.uuid4())

        try:
            doc = docx.Document(str(file_path))

            # Extract paragraphs
            paragraph_texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

            # Extract tables (vital for fee schedules, term tables, indemnities)
            table_texts = []
            for table in doc.tables:
                for row in table.rows:
                    row_content = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_content:
                        deduped = []
                        for cell_text in row_content:
                            if not deduped or deduped[-1] != cell_text:
                                deduped.append(cell_text)
                        table_texts.append(" | ".join(deduped))

            full_text = "\n\n".join(paragraph_texts + table_texts)
            if not full_text.strip():
                return [], 0, doc_id

            # Estimate pages: approx 2500 characters per standard contract page
            estimated_pages = max(1, (len(full_text) // 2500) + 1)
            raw_chunks = self.splitter.split_text(full_text)

            total_chunks = max(1, len(raw_chunks))
            for chunk_idx, chunk_text in enumerate(raw_chunks):
                page_num = min(estimated_pages, (chunk_idx * estimated_pages // total_chunks) + 1)
                chunks.append({
                    "chunk_id": f"{doc_id}_p{page_num}_c{chunk_idx}",
                    "doc_id": doc_id,
                    "filename": filename,
                    "page_number": page_num,
                    "text": chunk_text
                })

            return chunks, estimated_pages, doc_id
        except Exception as e:
            print(f"python-docx error processing {filename}: {e}")
            return [], 0, doc_id

    def process_document(self, file_path: Path, filename: str) -> Tuple[List[Dict[str, Any]], int, str]:
        """
        Unified document processor routing to Word (.docx) or PDF parser.
        """
        ext = Path(filename).suffix.lower()
        if ext == ".docx":
            return self.process_docx(file_path, filename)
        else:
            return self.process_pdf(file_path, filename)

pdf_service = PDFService()

